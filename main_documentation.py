import json
import re
from docx import Document
from docx.shared import Pt
import base64
from extraction import get_model_definition_lro, extract_tables_and_measures
from config import *
import re

DB_REGEX = re.compile(r'\[Name\s*=\s*"([^"]+)"\s*,\s*Kind\s*=\s*"Database"\]', re.IGNORECASE)
SCHEMA_REGEX = re.compile(r'\[Name\s*=\s*"([^"]+)"\s*,\s*Kind\s*=\s*"Schema"\]', re.IGNORECASE)
TABLE_REGEX = re.compile(r'\[Name\s*=\s*"([^"]+)"\s*,\s*Kind\s*=\s*"Table"\]', re.IGNORECASE)

def normalize_semantic_json(parsed_data: dict) -> dict:
    """
    Normalize extracted semantic model JSON into a stable schema
    consumed by Word documentation.
    """
    normalized = {}

    # --------------------------------------------------
    # Dashboard / Semantic model metadata
    # --------------------------------------------------
    # FIX: Use the correct key from parsed_data
    normalized["semantic_model_metadata"] = {
        "displayName": parsed_data.get("displayName", "Unknown Dashboard")
    }

    # --------------------------------------------------
    # Core model elements (pass-through)
    # --------------------------------------------------
    normalized["tables"] = parsed_data.get("tables", [])
    normalized["relationships"] = parsed_data.get("relationships", [])
    normalized["measures"] = parsed_data.get("measures", [])

    return normalized



class MainDocumentation:

    IGNORE_TABLE_PATTERNS = [
        r"^DateTableTemplate_",
        r"^LocalDateTable_",
        r"^measures$",
        r"^Measure$"
    ]

    def __init__(self, semantic_json_path, template_path, output_doc_path):
        self.semantic_json_path = semantic_json_path
        self.template_path = template_path
        self.output_doc_path = output_doc_path

         # Load template instead of empty doc
        self.document = Document(template_path)

        self.semantic_json = self._load_json()
#
    # --------------------------------------------------
    # Load JSON
    # --------------------------------------------------
    def _load_json(self):
        with open(self.semantic_json_path, "r", encoding="utf-8") as f:
            return json.load(f)
        

    def _find_placeholder_paragraph(self, placeholder):
        for p in self.document.paragraphs:
            if placeholder in p.text:
                return p
        return None


    # --------------------------------------------------
    # Ignore system tables
    # --------------------------------------------------
    def _should_ignore_table(self, table_name):
        for pattern in self.IGNORE_TABLE_PATTERNS:
            if re.match(pattern, table_name):
                return True
        return False

    # --------------------------------------------------
    # Build relationship lookup
    # --------------------------------------------------
    def _build_relationship_lookup(self):
        """
        Creates:
        {
        (TableName, ColumnName): (ConnectedTable, ConnectedColumn)
        }
        """
        lookup = {}

        for rel in self.semantic_json.get("relationships", []):
            ft = rel.get("fromTable")
            fc = rel.get("fromColumn")
            tt = rel.get("toTable")
            tc = rel.get("toColumn")

            if not all([ft, fc, tt, tc]):
                continue

            # 🚫 IGNORE LOCAL / SYSTEM TABLES
            if (
                self._should_ignore_table(ft)
                or self._should_ignore_table(tt)
            ):
                continue

            # forward
            lookup[(ft, fc)] = (tt, tc)
            # reverse (THIS enables Sheet1, Sheet1 (2), etc.)
            lookup[(tt, tc)] = (ft, fc)

        return lookup
    

    



    # --------------------------------------------------
    # Dashboard Name
    # --------------------------------------------------
    def add_dashboard_name(self):
        metadata = self.semantic_json.get("semantic_model_metadata", {})
        dashboard_name = metadata.get("displayName", "Unknown Dashboard")

        self.document.add_heading("Dashboard Overview", level=1)
        p = self.document.add_paragraph()
        r = p.add_run(f"Dashboard Name: {dashboard_name}")
        r.bold = True

    # --------------------------------------------------
    # Tables, Columns + Relationships (MAIN LOGIC)
    # --------------------------------------------------
    def add_tables_and_columns(self):
        tables = self.semantic_json.get("tables", [])
        relationship_lookup = self._build_relationship_lookup()

        self.document.add_heading("Data Model – Tables & Columns", level=1)

        for table in tables:
            table_name = table.get("name")

            if self._should_ignore_table(table_name):
                continue

            # Table heading
            self.document.add_heading(f"Table: {table_name}", level=2)

            # ADD SOURCE INFORMATION HERE
            source_info = table.get("source")
            if source_info:
                source_type = source_info.get("source_type", "Unknown")
                
                p = self.document.add_paragraph()
                p.add_run("Source Type: ").bold = True
                p.add_run(f"{source_type}\n")
                
                # Handle different source types
                if source_type == "Azure Databricks":
                    p.add_run("Database: ").bold = True
                    p.add_run(f"{source_info.get('database', 'N/A')}\n")
                    p.add_run("Data Product: ").bold = True
                    p.add_run(f"{source_info.get('data_product', 'N/A')}\n")
                    p.add_run("Table: ").bold = True
                    p.add_run(f"{source_info.get('table', 'N/A')}\n")
                
                elif "Excel" in source_type:
                    if source_info.get('file_name'):
                        p.add_run("File Name: ").bold = True
                        p.add_run(f"{source_info.get('file_name')}\n")
                    if source_info.get('url'):
                        p.add_run("URL: ").bold = True
                        p.add_run(f"{source_info.get('url')}\n")
                    elif source_info.get('file_path'):
                        p.add_run("File Path: ").bold = True
                        p.add_run(f"{source_info.get('file_path')}\n")
                
                elif source_type == "SQL Server":
                    if source_info.get('server'):
                        p.add_run("Server: ").bold = True
                        p.add_run(f"{source_info.get('server')}\n")
                    if source_info.get('database'):
                        p.add_run("Database: ").bold = True
                        p.add_run(f"{source_info.get('database')}\n")
                
                elif source_type == "SharePoint":
                    if source_info.get('url'):
                        p.add_run("URL: ").bold = True
                        p.add_run(f"{source_info.get('url')}\n")
                
                else:
                    # Unknown or other types
                    if source_info.get('datasource_expression'):
                        p.add_run("Source Expression: ").bold = True
                        # Truncate long expressions
                        expr = source_info.get('datasource_expression', '')
                        if len(expr) > 100:
                            expr = expr[:100] + "..."
                        p.add_run(f"{expr}\n")

            # Word table (6 columns)
            word_table = self.document.add_table(rows=1, cols=6)
            word_table.style = "Table Grid"

            header = word_table.rows[0].cells
            header[0].text = "Table Name"
            header[1].text = "Column Name"
            header[2].text = "Data Type"
            header[3].text = "Description"
            header[4].text = "Connects to Table"
            header[5].text = "Connects to Column"

            for column in table.get("columns", []):
                col_name = column.get("name", "")
                data_type = column.get("dataType", "Not Available")

                row = word_table.add_row().cells
                row[0].text = table_name
                row[1].text = col_name
                row[2].text = data_type
                row[3].text = ""

                # Relationship mapping
                rel = relationship_lookup.get((table_name, col_name))
                if rel:
                    row[4].text = rel[0]   # Connected Table
                    row[5].text = rel[1]   # Connected Column
                else:
                    row[4].text = ""
                    row[5].text = ""
            
            # Add spacing after each table
            self.document.add_paragraph()

    # --------------------------------------------------
    # Build Document
    # --------------------------------------------------
    # def build(self):
    #     self.add_dashboard_name()
    #     self.add_tables_and_columns()
    #     self.document.save(self.output_docx_path)
    #     print(f"✅ Word documentation created: {self.output_docx_path}")


    
    def build(self):
        # 1️⃣ Dashboard
        self.add_dashboard_name()

        # 2️⃣ Tables & Columns
        self.add_tables_and_columns()

        # 3️⃣ Measures (NEW – SAME DOC)
        measure_docs = MeasureDocumentation(
            semantic_json=self.semantic_json,
            document=self.document
        )
        measure_docs.add_measure_documentation()

        # 4️⃣ Save
        self.document.save(self.output_doc_path)
        print(f"✅ Word documentation created: {self.output_doc_path}")


    # --------------------------------------------------
    # Build Document
    # --------------------------------------------------
    ###############################################################################################################################
class TMDLSourceExtractor:
    """
    Extract table-level datasource information from TMDL table files.
    SIMPLE, SAFE, NO POLLING, NO OVER-REGEX.
    """

    IGNORE_TABLE_PREFIXES = (
        "DateTableTemplate_",
        "LocalDateTable_"
    )

    def __init__(self, model_definition):
        self.model_definition = model_definition

    # --------------------------------------------------
    # Databricks Data Product extractor (tbl_*)
    # --------------------------------------------------
    def _extract_databricks_dp(self, m_query: str):
        db_match = DB_REGEX.search(m_query)
        schema_match = SCHEMA_REGEX.search(m_query)
        table_match = TABLE_REGEX.search(m_query)

        if not (db_match and schema_match and table_match):
            return None

        return {
            "source_type": "Azure Databricks",
            "database": db_match.group(1),
            "data_product": schema_match.group(1),
            "table": table_match.group(1)
        }
    
    # --------------------------------------------------
    # Excel extractor
    # --------------------------------------------------
    def _extract_excel_source(self, source_expression: str):
        """Extract Excel file information"""
        # Check if it's an Excel source
        if "Excel.Workbook" not in source_expression:
            return None
        
        # Try to extract URL
        url_match = re.search(r'Web\.Contents\("([^"]+)"\)', source_expression)
        file_match = re.search(r'File\.Contents\("([^"]+)"\)', source_expression)
        
        if url_match:
            url = url_match.group(1)
            # Extract filename from URL
            filename = url.split('/')[-1] if '/' in url else url
            return {
                "source_type": "Excel (Web)",
                "file_name": filename,
                "url": url
            }
        elif file_match:
            filepath = file_match.group(1)
            filename = filepath.split('\\')[-1] if '\\' in filepath else filepath
            return {
                "source_type": "Excel (File)",
                "file_name": filename,
                "file_path": filepath
            }
        else:
            return {
                "source_type": "Excel",
                "source_details": "Excel Workbook"
            }
    
    # --------------------------------------------------
    # SQL Server extractor
    # --------------------------------------------------
    def _extract_sql_source(self, source_expression: str):
        """Extract SQL Server information"""
        if "Sql.Database" not in source_expression:
            return None
        
        server_match = re.search(r'"([^"]+)"', source_expression)
        db_match = re.search(r'"([^"]+)",\s*"([^"]+)"', source_expression)
        
        if db_match:
            return {
                "source_type": "SQL Server",
                "server": db_match.group(1),
                "database": db_match.group(2)
            }
        elif server_match:
            return {
                "source_type": "SQL Server",
                "server": server_match.group(1)
            }
        
        return None
    
    # --------------------------------------------------
    # SharePoint extractor
    # --------------------------------------------------
    def _extract_sharepoint_source(self, source_expression: str):
        """Extract SharePoint information"""
        if "SharePoint" not in source_expression:
            return None
        
        url_match = re.search(r'"(https?://[^"]+)"', source_expression)
        if url_match:
            return {
                "source_type": "SharePoint",
                "url": url_match.group(1)
            }
        
        return {
            "source_type": "SharePoint",
            "source_details": "SharePoint List/Library"
        }

    # --------------------------------------------------
    # Main extraction
    # --------------------------------------------------
    def extract_table_sources(self):
        table_sources = {}

        parts = self.model_definition.get("definition", {}).get("parts", [])

        for part in parts:
            path = part.get("path", "")
            payload = part.get("payload")

            # Only table TMDL files
            if not path.startswith("definition/tables/"):
                continue

            table_name = path.split("/")[-1].replace(".tmdl", "")

            if table_name.startswith(self.IGNORE_TABLE_PREFIXES):
                continue

            # Decode payload
            try:
                content = base64.b64decode(payload).decode("utf-8")
            except Exception:
                continue

            # Extract full Source block (not just first line)
            source_block_match = re.search(
                r"Source\s*=\s*([\s\S]+?)\n\s*in\s",
                content,
                re.IGNORECASE
            )

            source_block = source_block_match.group(1).strip() if source_block_match else ""

            # --------------------------------------------------
            # Try different extractors in priority order
            # --------------------------------------------------
            
            # 1. Databricks Data Product (tbl_*)
            if table_name.lower().startswith("tbl_") and source_block:
                dp_info = self._extract_databricks_dp(source_block)
                if dp_info:
                    table_sources[table_name] = dp_info
                    continue
            
            # 2. Excel
            excel_info = self._extract_excel_source(source_block)
            if excel_info:
                table_sources[table_name] = excel_info
                continue
            
            # 3. SQL Server
            sql_info = self._extract_sql_source(source_block)
            if sql_info:
                table_sources[table_name] = sql_info
                continue
            
            # 4. SharePoint
            sp_info = self._extract_sharepoint_source(source_block)
            if sp_info:
                table_sources[table_name] = sp_info
                continue

            # --------------------------------------------------
            # DEFAULT (existing behavior)
            # --------------------------------------------------
            source_line_match = re.search(
                r"Source\s*=\s*(.+)",
                content
            )

            source_expression = (
                source_line_match.group(1).strip()
                if source_line_match
                else "Not Found"
            )

            table_sources[table_name] = {
                "source_type": "Unknown",
                "datasource_expression": source_expression
            }

        return table_sources


###################################################################################################################################


COLUMN_REGEX = re.compile(
    r"'([^']+)'\[([^\]]+)\]|([A-Za-z0-9_]+)\[([^\]]+)\]"
)

MEASURE_REGEX = re.compile(
    r"\[(?![^\]]*\])(.*?)\]"
)

def extract_measure_dependencies(expression: str):
    tables = set()
    columns = set()
    measures = set()

    for match in COLUMN_REGEX.findall(expression):
        table = match[0] or match[2]
        column = match[1] or match[3]

        tables.add(table)
        columns.add(f"{table}[{column}]")

    for m in MEASURE_REGEX.findall(expression):
        if "[" not in m:
            measures.add(m)

    return {
        "source_tables": sorted(tables),
        "source_columns": sorted(columns),
        "referenced_measures": sorted(measures)
    }

class MeasureDocumentation:

    def __init__(self, semantic_json: dict, document: Document):
        self.semantic_json = semantic_json
        self.document = document

    # --------------------------------------------------
    # Collect measures from BOTH measure fields
    # --------------------------------------------------
    def _get_all_measures(self):
        measures = {}

        def normalize(name: str) -> str:
            # Normalize to avoid duplicates
            return (
                name.replace("\u00A0", " ")   # non-breaking space
                    .strip()
                    .lower()
            )

        # 1️⃣ Top-level measures
        for m in self.semantic_json.get("measures", []):
            name = m.get("name")
            expression = m.get("expression", "")

            if not name:
                continue

            if isinstance(expression, list):
                expression = "\n".join(expression)

            key = normalize(name)

            if key not in measures:
                measures[key] = {
                    "display_name": name.strip(),
                    "expression": expression
                }

        # 2️⃣ Measures inside tables
        for table in self.semantic_json.get("tables", []):
            for m in table.get("measures", []):
                name = m.get("name")
                expression = m.get("expression", "")

                if not name:
                    continue

                if isinstance(expression, list):
                    expression = "\n".join(expression)

                key = normalize(name)

                if key not in measures:
                    measures[key] = {
                        "display_name": name.strip(),
                        "expression": expression
                    }

        return measures



    # --------------------------------------------------
    # Add Measure Documentation Table
    # --------------------------------------------------
    def add_measure_documentation(self):

        measures = self._get_all_measures()

        self.document.add_heading("Measures – Source Mapping", level=1)

        word_table = self.document.add_table(rows=1, cols=5)
        word_table.style = "Table Grid"

        header = word_table.rows[0].cells
        header[0].text = "Measure Name"
        #header[1].text = "Source Table Names"
        header[2].text = "Source Columns"
        header[3].text = "Referenced Measures"
        header[4].text = "Description"

        for _, measure in measures.items():

            deps = extract_measure_dependencies(measure["expression"])

            row = word_table.add_row().cells
            row[0].text = measure["display_name"]
            #row[1].text = ", ".join(deps["source_tables"])
            row[2].text = ", ".join(deps["source_columns"])
            row[3].text = ", ".join(deps["referenced_measures"])
            row[4].text = ""




if __name__ == "__main__":

    # --------------------------------------------------
    # 1️⃣ Fetch model definition (Fabric API)
    # --------------------------------------------------
    model_definition = get_model_definition_lro(MODEL_ID)

    # --------------------------------------------------
    # 2️⃣ Extract tables, columns, measures, relationships
    # --------------------------------------------------
    parsed_data = extract_tables_and_measures(model_definition)
    
    # --------------------------------------------------
    # FIX: Extract displayName from model_definition
    # --------------------------------------------------
    # The displayName should be in the model definition response
    # Check different possible locations in the API response
    display_name = None
    
    # Try to find displayName in various locations
    if isinstance(model_definition, dict):
        # Option 1: Direct key
        display_name = model_definition.get("displayName")
        
        # Option 2: Inside definition
        if not display_name:
            definition = model_definition.get("definition", {})
            if isinstance(definition, dict):
                display_name = definition.get("displayName")
        
        # Option 3: Inside model metadata
        if not display_name:
            model = model_definition.get("model", {})
            if isinstance(model, dict):
                display_name = model.get("name") or model.get("displayName")
        
        # Option 4: Parse from TMDL definition
        if not display_name:
            parts = model_definition.get("definition", {}).get("parts", [])
            for part in parts:
                if part.get("path") == "definition/model.tmdl":
                    try:
                        payload = part.get("payload", "")
                        content = base64.b64decode(payload).decode("utf-8")
                        # Look for model name in TMDL
                        name_match = re.search(r'model\s+Model\s*\n.*?name:\s*(.+)', content, re.DOTALL)
                        if name_match:
                            display_name = name_match.group(1).strip().strip("'\"")
                    except:
                        pass
    
    # Use the found name or default
    parsed_data["displayName"] = display_name
    
    print(f"📊 Dashboard Name: {parsed_data['displayName']}")

    # --------------------------------------------------
    # 3️⃣ Normalize schema (🔥 CRITICAL STEP)
    # --------------------------------------------------
    semantic_json = normalize_semantic_json(parsed_data)

    # --------------------------------------------------
    # 4️⃣ Extract table sources from TMDL
    # --------------------------------------------------
    source_extractor = TMDLSourceExtractor(model_definition)
    table_sources = source_extractor.extract_table_sources()

    # --------------------------------------------------
    # 5️⃣ Attach sources to tables
    # --------------------------------------------------
    for table in semantic_json["tables"]:
        table_name = table.get("name")
        table["source"] = table_sources.get(table_name)

    # --------------------------------------------------
    # 6️⃣ Save FINAL semantic JSON
    # --------------------------------------------------
    OUTPUT_JSON = "4_semantic_model_with_table_sources.json"
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(semantic_json, f, indent=2)

    print("✅ Normalized semantic model JSON created")

    # --------------------------------------------------
    # 7️⃣ Build Word documentation
    # --------------------------------------------------
    docs = MainDocumentation(
        semantic_json_path=OUTPUT_JSON,
        output_docx_path="Doc_op/SIH_Dashboard_Documentation.docx"
    )
    docs.build()