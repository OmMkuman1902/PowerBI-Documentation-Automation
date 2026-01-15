import json
import re
from docx import Document

# --------------------------------------------------
# REGEX (FROM YOUR CODE)
# --------------------------------------------------

COLUMN_REGEX = re.compile(
    r"'([^']+)'\[([^\]]+)\]|([A-Za-z0-9_]+)\[([^\]]+)\]"
)

MEASURE_REGEX = re.compile(
    r"\[(?![^\]]*\])(.*?)\]"
)

def extract_measure_dependencies(expression: str):
    tables, columns, measures = set(), set(), set()

    for match in COLUMN_REGEX.findall(expression):
        table = match[0] or match[2]
        column = match[1] or match[3]
        tables.add(table)
        columns.add(f"{table}[{column}]")

    for m in MEASURE_REGEX.findall(expression):
        if "[" not in m:
            measures.add(m)

    return {
        "source_tables": sorted(tables),
        "source_columns": sorted(columns),
        "referenced_measures": sorted(measures)
    }

# --------------------------------------------------
# MAIN DOCUMENT CLASS
# --------------------------------------------------

class MainDocumentation:

    IGNORE_TABLE_PATTERNS = [
    r"^DateTableTemplate",
    r"^LocalDateTable",
    r"^measures?$",
    r"^Measure$"
]


    def __init__(self, semantic_json_path, template_path, output_doc_path):
        self.semantic_json_path = semantic_json_path
        self.template_path = template_path
        self.output_doc_path = output_doc_path

        self.document = Document(template_path)
        self.semantic_json = self._load_json()

    # --------------------------------------------------
    def _load_json(self):
        with open(self.semantic_json_path, "r", encoding="utf-8") as f:
            return json.load(f)

    # --------------------------------------------------
    def _should_ignore_table(self, table_name: str) -> bool:
        if not table_name:
            return True

        name = table_name.strip()

        for pattern in self.IGNORE_TABLE_PATTERNS:
            if re.match(pattern, name, re.IGNORECASE):
                return True

        return False


    # --------------------------------------------------
    def _build_relationship_lookup(self):
        lookup = {}

        for r in self.semantic_json.get("relationships", []):
            ft, fc = r.get("fromTable"), r.get("fromColumn")
            tt, tc = r.get("toTable"), r.get("toColumn")

            if not all([ft, fc, tt, tc]):
                continue

            if self._should_ignore_table(ft) or self._should_ignore_table(tt):
                continue

            lookup[(ft, fc)] = (tt, tc)
            lookup[(tt, tc)] = (ft, fc)

        return lookup

    # --------------------------------------------------
    def _find_placeholder(self, placeholder):

    # 1️⃣ Normal paragraphs
        for p in self.document.paragraphs:
            if placeholder in "".join(run.text for run in p.runs):
                return p

        # 2️⃣ Tables
        for table in self.document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if placeholder in "".join(run.text for run in p.runs):
                            return p

        # 3️⃣ Headers & Footers
        for section in self.document.sections:
            for hf in [section.header, section.footer]:
                for p in hf.paragraphs:
                    if placeholder in "".join(run.text for run in p.runs):
                        return p

        return None


    # ==================================================
    # 1️⃣ DASHBOARD NAME
    # ==================================================
    def insert_dashboard_name(self, placeholder="{{DASHBOARD_NAME}}"):

        p = self._find_placeholder(placeholder)
        if not p:
            raise ValueError("Dashboard name placeholder not found")

        name = self.semantic_json.get(
            "semantic_model_metadata", {}
        ).get("displayName", "Unknown Dashboard")

        p.text = f"Dashboard Name: {name}"

    # ==================================================
    # 2️⃣ DATA SOURCES TABLE
    # ==================================================
    def insert_data_sources(self, placeholder="{{DATA_SOURCES}}"):

        p = self._find_placeholder(placeholder)
        if not p:
            raise ValueError("Data sources placeholder not found")

        p.text = ""

        word_table = self.document.add_table(rows=1, cols=5)
        word_table.style = "Table Grid"

        hdr = word_table.rows[0].cells
        hdr[0].text = "Table Name"
        hdr[1].text = "Source Type"
        hdr[2].text = "Database / File"
        hdr[3].text = "Schema / Path"
        hdr[4].text = "Comments"

        for table in self.semantic_json.get("tables", []):
            table_name = table.get("name")
            source = table.get("source")

            if not source or self._should_ignore_table(table_name):
                continue

            row = word_table.add_row().cells
            row[0].text = table_name
            row[1].text = source.get("source_type", "Unknown")

            if source.get("source_type") == "Azure Databricks":
                row[2].text = source.get("database", "")
                row[3].text = source.get("data_product", "")
                row[4].text = source.get("table", "")

            elif "Excel" in source.get("source_type", ""):
                row[2].text = source.get("file_name", "")
                row[3].text = source.get("url") or source.get("file_path", "")
                row[4].text = ""

            elif source.get("source_type") == "SQL Server":
                row[2].text = source.get("database", "")
                row[3].text = source.get("server", "")
                row[4].text = ""

            elif source.get("source_type") == "SharePoint":
                row[2].text = ""
                row[3].text = source.get("url", "")
                row[4].text = ""

            else:
                row[2].text = ""
                row[3].text = ""
                row[4].text = "Unknown source"

        p._p.addnext(word_table._tbl)

    # ==================================================
    # 3️⃣ TABLES & COLUMNS (PER TABLE)
    # ==================================================
    def insert_tables_and_columns(self, placeholder="{{DATA_MODEL_TABLES}}"):

        p = self._find_placeholder(placeholder)
        if not p:
            raise ValueError("DATA_MODEL_TABLES placeholder not found")

        relationship_lookup = self._build_relationship_lookup()
        tables = self.semantic_json.get("tables", [])

        # Remove placeholder text
        p.text = ""

        for table in tables:
            table_name = table.get("name")

            if self._should_ignore_table(table_name):
                continue

            # -------------------------------
            # TABLE HEADING
            # -------------------------------
            heading = self.document.add_heading(f"Table: {table_name}", level=4)
            p._p.addnext(heading._p)

            # -------------------------------
            # SOURCE INFORMATION (NEW)
            # -------------------------------
            source = table.get("source")
            if source:
                src_p = self.document.add_paragraph()
                src_p.add_run("Source Type: ").bold = True
                src_p.add_run(source.get("source_type", "Unknown"))

                st = source.get("source_type")

                if st == "Azure Databricks":
                    src_p.add_run("\nDatabase: ").bold = True
                    src_p.add_run(source.get("database", ""))
                    src_p.add_run("\nData Product: ").bold = True
                    src_p.add_run(source.get("data_product", ""))
                    src_p.add_run("\nTable: ").bold = True
                    src_p.add_run(source.get("table", ""))

                elif "Excel" in st:
                    if source.get("file_name"):
                        src_p.add_run("\nFile Name: ").bold = True
                        src_p.add_run(source.get("file_name"))
                    if source.get("url"):
                        src_p.add_run("\nURL: ").bold = True
                        src_p.add_run(source.get("url"))
                    elif source.get("file_path"):
                        src_p.add_run("\nFile Path: ").bold = True
                        src_p.add_run(source.get("file_path"))

                elif st == "SQL Server":
                    if source.get("server"):
                        src_p.add_run("\nServer: ").bold = True
                        src_p.add_run(source.get("server"))
                    if source.get("database"):
                        src_p.add_run("\nDatabase: ").bold = True
                        src_p.add_run(source.get("database"))

                elif st == "SharePoint":
                    if source.get("url"):
                        src_p.add_run("\nURL: ").bold = True
                        src_p.add_run(source.get("url"))

                # Insert source paragraph under heading
                heading._p.addnext(src_p._p)

            # -------------------------------
            # COLUMNS & RELATIONSHIPS TABLE
            # -------------------------------
            word_table = self.document.add_table(rows=1, cols=6)
            word_table.style = "Table Grid"

            hdr = word_table.rows[0].cells
            hdr[0].text = "Table Name"
            hdr[1].text = "Column Name"
            hdr[2].text = "Data Type"
            hdr[3].text = "Description"
            hdr[4].text = "Connects to Table"
            hdr[5].text = "Connects to Column"

            for column in table.get("columns", []):
                col_name = column.get("name", "")
                data_type = column.get("dataType", "Not Available")

                row = word_table.add_row().cells
                row[0].text = table_name
                row[1].text = col_name
                row[2].text = data_type
                row[3].text = ""

                rel = relationship_lookup.get((table_name, col_name))
                if rel:
                    row[4].text = rel[0]
                    row[5].text = rel[1]
                else:
                    row[4].text = ""
                    row[5].text = ""

            # Insert table after source paragraph (or heading if no source)
            insert_after = src_p._p if source else heading._p
            insert_after.addnext(word_table._tbl)


    # ==================================================
    # 4️⃣ MEASURES
    # ==================================================
    def insert_measures(self, placeholder="{{MEASURES_MAPPING}}"):

        p = self._find_placeholder(placeholder)
        if not p:
            raise ValueError("Measures placeholder not found")

        p.text = ""

        wt = self.document.add_table(rows=1, cols=4)
        wt.style = "Table Grid"

        hdr = wt.rows[0].cells
        hdr[0].text = "Measure Name"
        hdr[1].text = "Source Columns"
        #hdr[2].text = "Referenced Measures"
        hdr[2].text = "Description"

        for m in self._get_all_measures().values():
            deps = extract_measure_dependencies(m["expression"])
            r = wt.add_row().cells
            r[0].text = m["display_name"]
            r[1].text = ", ".join(deps["source_columns"])
            #r[2].text = ", ".join(deps["referenced_measures"])
            r[2].text = ""

        p._p.addnext(wt._tbl)

    # --------------------------------------------------
    def _get_all_measures(self):
        measures = {}

        def norm(n): return n.strip().lower()

        for m in self.semantic_json.get("measures", []):
            measures[norm(m["name"])] = {
                "display_name": m["name"],
                "expression": "\n".join(m["expression"]) if isinstance(m["expression"], list) else m["expression"]
            }

        for t in self.semantic_json.get("tables", []):
            for m in t.get("measures", []):
                key = norm(m["name"])
                if key not in measures:
                    measures[key] = {
                        "display_name": m["name"],
                        "expression": "\n".join(m["expression"]) if isinstance(m["expression"], list) else m["expression"]
                    }

        return measures

    # --------------------------------------------------
    def build(self):
        self.insert_dashboard_name()
        self.insert_data_sources()
        self.insert_tables_and_columns()
        self.insert_measures()
        self.document.save(self.output_doc_path)
        print(f"✅ Document created: {self.output_doc_path}")

# --------------------------------------------------
# RUN TEST
# --------------------------------------------------

if __name__ == "__main__":

    docs = MainDocumentation(
        semantic_json_path=r"Doc_op\4_semantic_model_with_table_sources.json",
        template_path="Dashboarding - FDD Template.docx",
        output_path="FINAL_TEST_OUTPUT_scar.docx"
    )

    docs.build()
